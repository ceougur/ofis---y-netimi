#!/usr/bin/env python3
# Bağımsız mali model: Excel iş listesinden, programın kodunu KULLANMADAN, beklenen muhasebe sonuçları.
# Çıktı: islemler.json (yükleyicinin gireceği işlemler) ve beklenen.json (denetimde karşılaştırılacak sayılar).
# Kullanım: python3 model.py <excel> <cikti-klasoru>
import collections, datetime as dt, json, re, sys
from decimal import Decimal as D, ROUND_HALF_UP

import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
H = [c.value for c in wb["İş Listesi"][1]]
rows = [dict(zip(H, r)) for r in wb["İş Listesi"].iter_rows(min_row=2, values_only=True)]
DATE_TO = sys.argv[3] if len(sys.argv) > 3 else "9999-12-31"
_iso = lambda s: dt.datetime.strptime(s, "%d.%m.%Y").date().isoformat()
rows = [r for r in rows if r["İşlem Tipi"] == "Cari Açılış" or _iso(r["Tarih"]) <= DATE_TO]
cari_list = {r[0]: {"code": r[0], "name": r[1], "type": r[2]} for r in wb["Cari Listesi"].iter_rows(min_row=2, values_only=True)}
products = {r[0]: {"code": r[0], "name": r[1], "unit": r[2], "buy": r[3], "sell": r[4]} for r in wb["Ürün Listesi"].iter_rows(min_row=2, values_only=True)}

c2 = lambda x: D(str(x)).quantize(D("0.01"), ROUND_HALF_UP)
iso = lambda s: dt.datetime.strptime(s, "%d.%m.%Y").date().isoformat()
METHOD = {"Nakit": "cash", "Havale/EFT": "bank", "Kredi Kartı": "card"}
INSTR = {"Çek": "cheque", "Senet": "note"}


def add_months(day, n):
    d = dt.date.fromisoformat(day)
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    last = [31, 29 if y % 4 == 0 and (y % 100 or y % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return dt.date(y, m, min(d.day, last)).isoformat()


def line_money(qty, price, rate, included=False):
    base = c2(D(str(qty)) * D(str(price)))
    if included:
        net = (base * 100 / (100 + rate)).quantize(D("0.01"), ROUND_HALF_UP)
        vat = base - net
    else:
        net = base
        vat = (net * rate / 100).quantize(D("0.01"), ROUND_HALF_UP)
    return net, vat, net + vat


GIDER_KOD = {"Kira": "rent", "Elektrik": "utilities", "Su": "utilities", "Doğalgaz": "utilities", "İnternet": "telecom", "Telefon": "telecom", "Mobil": "telecom",
             "Kırtasiye": "office", "Yakıt": "fuel", "Akaryakıt": "fuel", "Bakım": "repair", "Onarım": "repair", "Danışmanlık": "advisory", "Muhasebe": "advisory", "Mali Müşavir": "advisory",
             "Yemek": "food", "Sigorta": "insurance", "Yazılım": "software", "Temizlik": "cleaning", "Güvenlik": "cleaning", "Reklam": "marketing", "Kargo": "freight", "Nakliye": "freight"}


def gider_kodu(text):
    for key, code in GIDER_KOD.items():
        if key.lower() in (text or "").lower():
            return code
    return "other"


ops = []
findings = collections.defaultdict(list)
E = collections.defaultdict(lambda: D(0))  # toplam sayaçlar
cari = collections.defaultdict(lambda: D(0))  # + bize borçlu, − biz borçluyuz
method_flow = collections.defaultdict(lambda: D(0))  # "cash:in" ...
stock = collections.defaultdict(lambda: D(0))
cheq = {"in": [D(0), 0], "out": [D(0), 0]}
plans = {}
cari_open = {}
excel_diff = D(0)

for r in rows:
    t = r["İşlem Tipi"]
    if t == "Cari Açılış":
        cl = cari_list.get(r["Cari Kodu"])
        cari_open[r["Cari Kodu"]] = iso(r["Tarih"])
        ops.append({"no": r["İşlem No"], "type": "cari", "date": iso(r["Tarih"]), "code": r["Cari Kodu"], "name": r["Cari Ünvanı"], "kind": {"Müşteri": "customer", "Tedarikçi": "supplier"}.get(cl["type"] if cl else "", "customer"), "excelType": cl["type"] if cl else ""})

used_cari = {r["Cari Kodu"] for r in rows if r["Cari Kodu"] not in ("-",)}
for code in sorted(used_cari - set(cari_open)):
    findings["cari_listesinde_yok"].append(code)

for r in rows:
    t, no, day = r["İşlem Tipi"], r["İşlem No"], iso(r["Tarih"])
    pay, status = r["Ödeme Şekli"], r["Durum"]
    code = r["Cari Kodu"]
    if t == "Cari Açılış":
        continue
    if code in cari_open and day < cari_open[code]:
        findings["cari_acilisindan_once"].append(no)
    if t in ("Mal Satışı", "Mal Alımı"):
        sale = t == "Mal Satışı"
        rate = int(r["KDV Oranı %"])
        net, vat, gross = line_money(r["Miktar"], r["Birim Fiyat (TL)"], rate)
        excel_diff += abs(net - c2(r["Toplam Tutar (TL)"]))
        due = iso(r["Vade Tarihi"]) if r["Vade Tarihi"] not in ("-", None) else day
        op = {"no": no, "type": "sale" if sale else "purchase", "date": day, "code": code, "product": r["Ürün/Hizmet Kodu"], "qty": r["Miktar"], "price": r["Birim Fiyat (TL)"], "vatRate": rate,
              "net": str(net), "vat": str(vat), "payable": str(gross), "excelNet": str(r["Toplam Tutar (TL)"]), "excelPayable": str(r["Genel Toplam (TL)"]), "payMethod": pay, "status": status, "due": due}
        stock[r["Ürün/Hizmet Kodu"]] += D(str(r["Miktar"])) * (-1 if sale else 1)
        E["sale_net" if sale else "purchase_net"] += net
        E["sale_vat" if sale else "purchase_vat"] += vat
        E["sale_payable" if sale else "purchase_payable"] += gross
        cari[code] += gross if sale else -gross
        if pay == "Taksitli":
            n = int(r["Taksit Sayısı"])
            op.update(mode="installments", count=n, firstDue=add_months(day, 1))
            plans[no] = {"total": gross, "paid": D(0), "count": n, "code": code}
            if status != "Taksitli Devam Ediyor":
                findings[f"taksitli_{status}_tahsilat_satiri_yok"].append(no)
        elif status == "Tamamlandı" and pay in METHOD:
            op.update(mode="paid", method=METHOD[pay])
            method_flow[f"{METHOD[pay]}:{'in' if sale else 'out'}"] += gross
            cari[code] += -gross if sale else gross
        elif status == "Tamamlandı" and pay in INSTR:
            op.update(mode="cheque", instrument=INSTR[pay])
            cheq["in" if sale else "out"][0] += gross
            cheq["in" if sale else "out"][1] += 1
            cari[code] += -gross if sale else gross
        else:
            op.update(mode="open")
            E["open_" + ("sale" if sale else "purchase")] += gross
            if status == "Kısmi Ödendi":
                findings["kismi_odendi_tutar_yok"].append(no)
        ops.append(op)
    elif t in ("Tahsilat", "Ödeme"):
        inflow = t == "Tahsilat"
        amount = c2(r["Genel Toplam (TL)"])
        due = iso(r["Vade Tarihi"]) if r["Vade Tarihi"] not in ("-", None) else day
        m = re.match(r"(ISL-\d+) nolu taksitli satışın (\d+)\. taksit", r["Açıklama"] or "")
        op = {"no": no, "type": "collect" if inflow else "pay", "date": day, "code": code, "amount": str(amount), "payMethod": pay, "due": due}
        if m:
            op.update(type="installment", sale=m.group(1), seq=int(m.group(2)))
            plans[m.group(1)]["paid"] += amount
        if pay in METHOD:
            op["method"] = METHOD[pay]
            method_flow[f"{METHOD[pay]}:{'in' if inflow else 'out'}"] += amount
        else:
            op["instrument"] = INSTR[pay]
            cheq["in" if inflow else "out"][0] += amount
            cheq["in" if inflow else "out"][1] += 1
        cari[code] += -amount if inflow else amount
        E["collect" if inflow else "pay"] += amount
        ops.append(op)
    elif t == "Gider Ödemesi":
        rate = int(r["KDV Oranı %"])
        gross = c2(r["Genel Toplam (TL)"])
        net, vat, _ = line_money(1, gross, rate, included=True)
        if rate and c2(r["KDV Tutarı (TL)"]) != vat:
            findings["gider_kdv_farki"].append(no)
        ops.append({"no": no, "type": "expense", "date": day, "category": r["Cari Ünvanı"], "itemName": r["Ürün/Hizmet Adı"], "expenseCode": gider_kodu(f"{r['Cari Ünvanı']} {r['Ürün/Hizmet Adı']}"),
                    "gross": str(gross), "net": str(net), "vat": str(vat), "vatRate": rate, "method": METHOD[pay]})
        method_flow[f"{METHOD[pay]}:out"] += gross
        E["expense_gross"] += gross
        E["expense_net"] += net
        E["expense_vat"] += vat
    elif t == "Maaş Ödemesi":
        amount = c2(r["Genel Toplam (TL)"])
        ops.append({"no": no, "type": "salary", "date": day, "code": code, "name": r["Cari Ünvanı"], "amount": str(amount), "method": METHOD[pay], "note": r["Açıklama"]})
        method_flow[f"{METHOD[pay]}:out"] += amount
        E["salary"] += amount
    else:
        raise SystemExit(f"bilinmeyen işlem tipi {t}")

personnel = sorted({(r["Cari Kodu"], r["Cari Ünvanı"]) for r in rows if r["İşlem Tipi"] == "Maaş Ödemesi"})
categories = sorted({r["Cari Ünvanı"] for r in rows if r["İşlem Tipi"] == "Gider Ödemesi"})
over = [k for k, p in plans.items() if p["paid"] > p["total"]]
expected = {
    "counts": dict(collections.Counter(r["İşlem Tipi"] for r in rows)),
    "money": {k: str(v) for k, v in E.items()},
    "kdv": {"hesaplanan": str(E["sale_vat"]), "indirilecek": str(E["purchase_vat"] + E["expense_vat"]), "odenecek": str(E["sale_vat"] - E["purchase_vat"] - E["expense_vat"])},
    "methods": {k: str(v) for k, v in sorted(method_flow.items())},
    "balances": {m: str(method_flow[f"{m}:in"] - method_flow[f"{m}:out"]) for m in ("cash", "bank", "card")},
    "cari": {k: str(v) for k, v in sorted(cari.items())},
    "cari_totals": {"borclu": str(sum(v for v in cari.values() if v > 0)), "alacakli": str(sum(-v for v in cari.values() if v < 0)), "net": str(sum(cari.values()))},
    "stock": {k: str(v) for k, v in sorted(stock.items())},
    "cheques": {"in_total": str(cheq["in"][0]), "in_count": cheq["in"][1], "out_total": str(cheq["out"][0]), "out_count": cheq["out"][1]},
    "plans": {"count": len(plans), "total": str(sum(p["total"] for p in plans.values())), "paid": str(sum(p["paid"] for p in plans.values())), "remaining": str(sum(p["total"] - p["paid"] for p in plans.values()))},
    "excel_arithmetic_diff": str(excel_diff),
    "findings": {k: {"count": len(v), "sample": v[:8]} for k, v in findings.items()},
    "plans_overpaid": over,
}
suffix = "" if DATE_TO == "9999-12-31" else f"-{DATE_TO}"
if suffix:
    json.dump(expected, open(f"{out}/beklenen{suffix}.json", "w"), ensure_ascii=False, indent=1)
    print("yalnız", f"beklenen{suffix}.json"); sys.exit(0)
json.dump({"ops": ops, "cari": list(cari_list.values()), "products": list(products.values()), "personnel": personnel, "categories": categories}, open(f"{out}/islemler.json", "w"), ensure_ascii=False, indent=0)
json.dump(expected, open(f"{out}/beklenen.json", "w"), ensure_ascii=False, indent=1)
print(json.dumps({k: expected[k] for k in ("counts", "money", "kdv", "balances", "cari_totals", "cheques", "plans", "excel_arithmetic_diff", "findings")}, ensure_ascii=False, indent=1))
print("stok eksi kalan ürün:", sum(1 for v in stock.values() if v < 0), "ops:", len(ops))
